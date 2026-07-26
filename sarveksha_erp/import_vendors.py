import frappe
import openpyxl
import os

def execute():
    file_path = os.path.abspath(os.path.join(frappe.get_app_path("sarveksha_erp"), "..", "datasets-needed", "Updated-vendorlist.xlsx"))
    
    if not os.path.exists(file_path):
        print(f"Error: Could not find {file_path}")
        return
        
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    
    # Map headers to indices
    try:
        col_company = headers.index("Company")
        col_pan = headers.index("PAN")
        col_gstin = headers.index("GSTIN")
        col_bank = headers.index("Bank Name/Branch")
        col_acc = headers.index("Account No.")
        col_ifsc = headers.index("IFSC")
    except ValueError as e:
        print(f"Error finding columns in Excel: {e}")
        return

    count = 0
    # Start from row 2
    for row in ws.iter_rows(min_row=2, values_only=True):
        company_name = row[col_company]
        
        if not company_name:
            continue
            
        # Clean data
        pan = str(row[col_pan] or "").strip()
        gstin = str(row[col_gstin] or "").strip()
        bank = str(row[col_bank] or "").strip()
        acc = str(row[col_acc] or "").strip()
        ifsc = str(row[col_ifsc] or "").strip()

        # Check if supplier exists
        if frappe.db.exists("Supplier", company_name):
            doc = frappe.get_doc("Supplier", company_name)
            doc.custom_pan = pan
            doc.tax_id = gstin
            doc.custom_bank_name = bank
            doc.custom_account_number = acc
            doc.custom_ifsc = ifsc
            doc.save(ignore_permissions=True)
            print(f"Updated existing supplier: {company_name}")
        else:
            doc = frappe.get_doc({
                "doctype": "Supplier",
                "supplier_name": company_name,
                "supplier_group": "All Supplier Groups",
                "supplier_type": "Company",
                "custom_pan": pan,
                "tax_id": gstin,
                "custom_bank_name": bank,
                "custom_account_number": acc,
                "custom_ifsc": ifsc,
                "custom_gst_applicable": 1 if gstin else 0,
            })
            doc.insert(ignore_permissions=True)
            print(f"Created new supplier: {company_name}")
            
        count += 1
        
    frappe.db.commit()
    print(f"\nSuccessfully imported/updated {count} vendors!")
