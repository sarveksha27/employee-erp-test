import frappe
from openpyxl import load_workbook
import os

def execute():
    frappe.set_user("Administrator")
    print("=" * 60)
    print("STARTING DATA SETUP...")
    print("=" * 60)

    # 1. CLEANUP OBSOLETE COMPANIES
    if frappe.db.exists("Company", "Baani Minerals"):
        frappe.delete_doc("Company", "Baani Minerals", force=True)
        print("  ✓ Removed obsolete company: Baani Minerals")

    # 2. CREATE & UPDATE OFFICIAL COMPANIES
    company_records = [
        {
            "company_name": "Sarveksha Realty and Inframine LLP",
            "abbr": "SRIL",
            "country": "India",
            "default_currency": "INR",
            "tax_id": "27AFTFS2557J1Z8",
            "custom_registration_no": "ACS-0698",
            "custom_import_export_code": "AFTFS2557J",
            "custom_default_port": "Mundra, JNPT",
            "phone_no": "+91 9769008220",
            "email": "sudheerg@sarveksha.com",
            "website": "www.sarveksha.com"
        },
        {
            "company_name": "Sarveksha BSTP SAS",
            "abbr": "SBSTP",
            "country": "Guinea",
            "default_currency": "USD",
            "tax_id": "684060320",
            "custom_registration_no": "RCCM / GN / TCC .2024. B. 02459",
            "custom_default_port": "Conakry",
            "phone_no": "+224 626407133",
            "email": "svbstp@gmail.com",
            "website": "www.sarveksha.com"
        },
        {
            "company_name": "Sarveksha Mining SARL",
            "abbr": "SMS",
            "country": "Cameroon",
            "default_currency": "USD",
            "custom_registration_no": "CM-NSI-02-2025-B12-00560",
            "custom_default_port": "Douala",
            "phone_no": "+237652280756",
            "email": "cameroonprojects@sarveksha.com",
            "website": "www.sarveksha.com"
        },
        {
            "company_name": "Baani Resources SARL",
            "abbr": "BRS",
            "country": "Cameroon",
            "default_currency": "USD",
            "custom_registration_no": "CM-NSI-02-2025-B12-00784",
            "custom_default_port": "Douala",
            "phone_no": "+237652280756",
            "email": "cameroonprojects@sarveksha.com",
            "website": "www.sarveksha.com"
        },
        {
            "company_name": "Sarveksha Botswana Proprietary Limited",
            "abbr": "SBPL",
            "country": "Botswana",
            "default_currency": "USD",
            "custom_registration_no": "BW00009608412",
            "custom_default_port": "Durban",
            "phone_no": "+237652280756",
            "email": "sudheerg@sarveksha.com",
            "website": "www.sarveksha.com"
        },
        {
            "company_name": "Sarveksha SL Limited",
            "abbr": "SSL",
            "country": "Sierra Leone",
            "default_currency": "USD",
            "tax_id": "1001412648",
            "custom_registration_no": "SLI 20624 SARVE22162",
            "custom_default_port": "Freetown",
            "phone_no": "+232 74899999",
            "email": "svbstp@gmail.com",
            "website": "www.sarveksha.com"
        },
        {
            "company_name": "Odhav Holdings",
            "abbr": "OH",
            "country": "Mauritius",
            "default_currency": "USD",
            "email": "svbstp@gmail.com",
            "website": "www.sarveksha.com"
        },
        {
            "company_name": "Globe Multitrade and Service LLC",
            "abbr": "GMT",
            "country": "United Arab Emirates",
            "default_currency": "USD",
            "website": "www.sarveksha.com"
        }
    ]

    for cdata in company_records:
        cname = cdata["company_name"]
        if not frappe.db.exists("Company", cname):
            doc = frappe.new_doc("Company")
            doc.update(cdata)
            doc.insert(ignore_permissions=True)
            print(f"  ✓ Created Company: {cname}")
        else:
            for k, v in cdata.items():
                if k != "company_name":
                    frappe.db.set_value("Company", cname, k, v)
            print(f"  ✓ Updated Company: {cname}")

    frappe.db.commit()

    # 3. IMPORT VENDOR EXCEL DATA (WITH BANK & CONTACT DETAILS)
    app_path = frappe.get_app_path("sarveksha_erp")
    file_sarveksha = os.path.join(app_path, "..", "datasets-needed", "VendorList-sarveksha.xlsx")
    file_updated = os.path.join(app_path, "..", "datasets-needed", "Updated-vendorlist.xlsx")

    vendor_data_map = {}

    if os.path.exists(file_sarveksha):
        try:
            wb1 = load_workbook(file_sarveksha, data_only=True)
            sheet1 = wb1.active
            for row in sheet1.iter_rows(min_row=2, values_only=True):
                if not row or not row[1]: continue
                s_name = str(row[1]).strip()
                if s_name.lower().startswith("same as"): continue
                vendor_data_map[s_name] = {
                    "supplier_name": s_name,
                    "custom_pan": str(row[2]).strip() if row[2] else None,
                    "custom_gstin": str(row[3]).strip() if row[3] else None,
                    "custom_bank_name": str(row[4]).strip() if row[4] else None,
                    "custom_bank_branch": str(row[5]).strip() if row[5] else None,
                    "custom_bank_account_no": str(row[6]).strip() if row[6] else None,
                    "custom_ifsc": str(row[7]).strip() if row[7] else None,
                    "custom_contact_email": str(row[8]).strip() if row[8] else None,
                    "custom_contact_phone": str(row[9]).strip() if row[9] else None,
                    "custom_contact_person": str(row[10]).strip() if row[10] else None,
                }
        except Exception as e:
            print(f"Error reading VendorList-sarveksha.xlsx: {e}")

    if os.path.exists(file_updated):
        try:
            wb2 = load_workbook(file_updated, data_only=True)
            sheet2 = wb2.active
            for row in sheet2.iter_rows(min_row=2, values_only=True):
                if not row or not row[1]: continue
                s_name = str(row[1]).strip()
                if s_name.lower().startswith("same as") or "same as" in str(row[2] or "").lower(): continue
                entry = vendor_data_map.get(s_name, {"supplier_name": s_name})
                if row[2]: entry["custom_pan"] = str(row[2]).strip()
                if row[3]: entry["custom_gstin"] = str(row[3]).strip()
                if row[4]: entry["custom_bank_name"] = str(row[4]).strip()
                if len(row) > 5 and row[5]: entry["custom_bank_account_no"] = str(row[5]).strip()
                if len(row) > 6 and row[6]: entry["custom_ifsc"] = str(row[6]).strip()
                if len(row) > 7 and row[7]: entry["custom_contact_email"] = str(row[7]).strip()
                if len(row) > 8 and row[8]: entry["custom_contact_phone"] = str(row[8]).strip()
                if len(row) > 9 and row[9]: entry["custom_contact_person"] = str(row[9]).strip()
                vendor_data_map[s_name] = entry
        except Exception as e:
            print(f"Error reading Updated-vendorlist.xlsx: {e}")

    created_count, updated_count = 0, 0
    for s_name, data in vendor_data_map.items():
        existing_name = s_name if frappe.db.exists("Supplier", s_name) else None
        if not existing_name:
            matches = frappe.get_all("Supplier", filters={"supplier_name": ["like", f"%{s_name}%"]}, fields=["name"])
            if matches: existing_name = matches[0].name

        fields_to_update = {
            "supplier_group": "All Supplier Groups",
            "tax_id": data.get("custom_gstin") or data.get("custom_pan"),
            "custom_pan": data.get("custom_pan"),
            "custom_gstin": data.get("custom_gstin"),
            "custom_bank_name": data.get("custom_bank_name"),
            "custom_bank_branch": data.get("custom_bank_branch"),
            "custom_bank_account_no": data.get("custom_bank_account_no"),
            "custom_ifsc": data.get("custom_ifsc"),
            "custom_contact_email": data.get("custom_contact_email"),
            "custom_contact_phone": data.get("custom_contact_phone"),
            "custom_contact_person": data.get("custom_contact_person"),
        }
        fields_to_update = {k: v for k, v in fields_to_update.items() if v is not None}

        if existing_name:
            for k, v in fields_to_update.items():
                frappe.db.set_value("Supplier", existing_name, k, v)
            updated_count += 1
        else:
            supplier = frappe.new_doc("Supplier")
            supplier.supplier_name = s_name
            supplier.update(fields_to_update)
            supplier.insert(ignore_permissions=True)
            created_count += 1

    frappe.db.commit()
    print(f"  ✓ Processed Suppliers: {created_count} created, {updated_count} updated.")

    print("\n✅ DATA SETUP COMPLETE.")