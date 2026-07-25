import frappe
from openpyxl import load_workbook
import os

def execute():
    frappe.set_user("Administrator")
    print("=" * 60)
    print("FIXING VENDORS AND DATA")
    print("=" * 60)

    # 1. Clear Global Default Company
    gd = frappe.get_single("Global Defaults")
    if gd.default_company == "Test Enterprise":
        gd.default_company = ""
        gd.save(ignore_permissions=True)
        print("✓ Cleared Test Enterprise from Global Defaults")

    # 2. Delete Suppliers named 1, 2, 3...
    suppliers = frappe.get_all("Supplier", fields=["name"])
    deleted = 0
    for s in suppliers:
        if s.name.isdigit():
            frappe.delete_doc("Supplier", s.name, ignore_permissions=True, force=True)
            deleted += 1
    print(f"✓ Deleted {deleted} incorrectly named suppliers.")

    # 3. Create Custom Fields on Supplier
    custom_fields = [
        {"fieldname": "custom_pan", "label": "PAN", "fieldtype": "Data", "insert_after": "tax_id"},
        {"fieldname": "custom_gstin", "label": "GSTIN", "fieldtype": "Data", "insert_after": "custom_pan"},
        {"fieldname": "custom_bank_name", "label": "Bank Name", "fieldtype": "Data", "insert_after": "custom_gstin"},
        {"fieldname": "custom_bank_branch", "label": "Bank Branch", "fieldtype": "Data", "insert_after": "custom_bank_name"},
        {"fieldname": "custom_bank_account_no", "label": "Account No.", "fieldtype": "Data", "insert_after": "custom_bank_branch"},
        {"fieldname": "custom_ifsc", "label": "IFSC", "fieldtype": "Data", "insert_after": "custom_bank_account_no"},
        {"fieldname": "custom_contact_email", "label": "Contact Email", "fieldtype": "Data", "insert_after": "custom_ifsc"},
        {"fieldname": "custom_contact_phone", "label": "Contact Phone", "fieldtype": "Data", "insert_after": "custom_contact_email"},
        {"fieldname": "custom_contact_person", "label": "Contact Person", "fieldtype": "Data", "insert_after": "custom_contact_phone"},
    ]

    for cf in custom_fields:
        if not frappe.db.exists("Custom Field", {"dt": "Supplier", "fieldname": cf["fieldname"]}):
            doc = frappe.new_doc("Custom Field")
            doc.dt = "Supplier"
            doc.fieldname = cf["fieldname"]
            doc.label = cf["label"]
            doc.fieldtype = cf["fieldtype"]
            doc.insert_after = cf["insert_after"]
            doc.insert(ignore_permissions=True)
            print(f"  ✓ Created Custom Field: {cf['label']}")
    
    frappe.db.commit()

    # 4. Import Excel Data
    file_path = os.path.join(frappe.get_app_path("sarveksha_erp"), "..", "datasets-needed", "VendorList-sarveksha.xlsx")
    if os.path.exists(file_path):
        print(f"\nProcessing Vendor Data from {file_path}...")
        try:
            wb = load_workbook(file_path, data_only=True)
            sheet = wb.active
            
            headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]
            print(f"Found headers: {headers}")
            
            # Find indices
            h_idx = {h: i for i, h in enumerate(headers)}
            
            name_idx = h_idx.get("Company", 1)  # The column is called 'Company'
            pan_idx = h_idx.get("PAN")
            gstin_idx = h_idx.get("GSTIN")
            b_name_idx = h_idx.get("Bank Name")
            b_branch_idx = h_idx.get("Bank Branch")
            b_acc_idx = h_idx.get("Account No.")
            ifsc_idx = h_idx.get("IFSC")
            email_idx = h_idx.get("Contact Email")
            phone_idx = h_idx.get("Contact Phone")
            person_idx = h_idx.get("Contact person")
                
            created_count = 0
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                name = row[name_idx].value
                if not name:
                    continue
                name = str(name).strip()
                
                group = "All Supplier Groups"
                
                if not frappe.db.exists("Supplier", name):
                    supplier = frappe.new_doc("Supplier")
                    supplier.supplier_name = name
                    supplier.supplier_group = group
                    
                    if pan_idx is not None and row[pan_idx].value: supplier.custom_pan = str(row[pan_idx].value)
                    if gstin_idx is not None and row[gstin_idx].value: supplier.custom_gstin = str(row[gstin_idx].value)
                    if b_name_idx is not None and row[b_name_idx].value: supplier.custom_bank_name = str(row[b_name_idx].value)
                    if b_branch_idx is not None and row[b_branch_idx].value: supplier.custom_bank_branch = str(row[b_branch_idx].value)
                    if b_acc_idx is not None and row[b_acc_idx].value: supplier.custom_bank_account_no = str(row[b_acc_idx].value)
                    if ifsc_idx is not None and row[ifsc_idx].value: supplier.custom_ifsc = str(row[ifsc_idx].value)
                    if email_idx is not None and row[email_idx].value: supplier.custom_contact_email = str(row[email_idx].value)
                    if phone_idx is not None and row[phone_idx].value: supplier.custom_contact_phone = str(row[phone_idx].value)
                    if person_idx is not None and row[person_idx].value: supplier.custom_contact_person = str(row[person_idx].value)
                    
                    supplier.insert(ignore_permissions=True)
                    created_count += 1
                    
            frappe.db.commit()
            print(f"  ✓ Successfully imported {created_count} Suppliers with all data.")
        except Exception as e:
            print(f"Error parsing Excel file: {e}")
    else:
        print(f"Excel file not found at {file_path}")

    print("\n✅ DATA FIX COMPLETE.")
